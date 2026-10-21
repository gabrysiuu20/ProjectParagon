from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from backend.calculator import calculate


def safe_text(cell, value):
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.data_type = "s"  # OCR and person names must never become spreadsheet formulas.


def export_xlsx(receipt, settlement, people):
    result = calculate(receipt, settlement, {p["id"] for p in people})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Produkty"
    sheet.append(["Produkt", "Ilość", "Cena jednostkowa", "Wartość", "Rabat", "Po rabacie", "Osoby"])
    assignments = {a.item_index: a.person_ids for a in settlement.assignments}
    names = {p["id"]: p["name"] for p in people}
    for index, item in enumerate(receipt.items):
        sheet.append([None, item.quantity, item.unit_price, item.amount, item.discount, item.amount-item.discount, None])
        safe_text(sheet.cell(index+2, 1), item.name)
        safe_text(sheet.cell(index+2, 7), ", ".join(names[p] for p in assignments.get(index, [])))
        for column in range(3, 7):
            sheet.cell(index+2, column).number_format = '0.00'
    sheet = workbook.create_sheet("Rozliczenie")
    sheet.append(["Osoba", "Należność PLN"])
    for person_id, amount in result["totals"].items():
        sheet.append([None, amount])
        safe_text(sheet.cell(sheet.max_row, 1), names[person_id])
    sheet.append(["Nieprzypisane", result["unassigned_amount"]])
    sheet.append(["Suma paragonu", receipt.total])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
